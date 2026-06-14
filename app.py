from flask import Flask, render_template, request, send_file
from PIL import Image
from pdf2image import convert_from_path
import os
from pptx import Presentation
import json
import pandas as pd
from fpdf import FPDF
import cairosvg
from PIL import Image
import requests
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
import qrcode
import rarfile
import zipfile
from werkzeug.utils import secure_filename
import py7zr
from reportlab.pdfgen import canvas
import subprocess
import uuid
from PyPDF2 import PdfMerger, PdfReader, PdfWriter
import pikepdf
from docx2pdf import convert

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
CONVERTED_FOLDER = "converted"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(CONVERTED_FOLDER, exist_ok=True)

@app.route("/")
def home():
    return render_template("index.html")
@app.route("/test")
def test():
    return "TEST OK"

@app.route("/convert", methods=["POST"])
def convert():

    file = request.files["image"]
    target_format = request.form["format"]

    if file.filename == "":
        return "No file selected"

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    img = Image.open(input_path)

    if img.mode == "RGBA":
        img = img.convert("RGB")

    output_file = f"converted.{target_format.lower()}"

    output_path = os.path.join(
        CONVERTED_FOLDER,
        output_file
    )

    if target_format == "JPG":
        img.save(output_path, "JPEG")
    else:
        img.save(output_path, target_format)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/image-converter")
def image_converter():
    return render_template("image-converter.html")
@app.route("/pdf-converter")
def pdf_converter():
    return render_template("pdf-converter.html")
@app.route("/convert-pdf", methods=["POST"])
def convert_pdf():

    file = request.files["pdf"]

    if file.filename == "":
        return "No PDF selected"

    pdf_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(pdf_path)

    pages = convert_from_path(pdf_path)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "pdf-page.jpg"
    )

    pages[0].save(output_path, "JPEG")

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/word-to-pdf", methods=["POST"])
def word_to_pdf():

    file = request.files["file"]

    docx_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(docx_path)

    pdf_path = os.path.join(
        CONVERTED_FOLDER,
        "converted.pdf"
    )

    convert(docx_path, pdf_path)

    return send_file(
        pdf_path,
        as_attachment=True
    )


@app.route("/pdf-split")
def pdf_split_page():
    return render_template("pdf-split.html")


    file.save(docx_path)

    pdf_path = os.path.join(
        CONVERTED_FOLDER,
        "converted.pdf"
    )

    convert(docx_path, pdf_path)

    return send_file(
        pdf_path,
        as_attachment=True
    )


@app.route("/image-to-pdf", methods=["POST"])
def image_to_pdf():

    file = request.files["image"]

    image_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(image_path)

    image = Image.open(image_path)


    if image.mode == "RGBA":
        image = image.convert("RGB")

    pdf_path = os.path.join(
        CONVERTED_FOLDER,
        "image.pdf"
    )

    image.save(pdf_path, "PDF")

    return send_file(
        pdf_path,
        as_attachment=True
    )
@app.route("/word-converter")
def word_converter():
    return render_template("word-converter.html")


@app.route("/image-to-pdf-converter")
def image_to_pdf_converter():
    return render_template("image-to-pdf.html")



@app.route("/about")
def about():
    return render_template("about.html")


@app.route("/privacy")
def privacy():
    return render_template("privacy.html")


@app.route("/contact")
def contact():
    return render_template("contact.html")


@app.route("/disclaimer")
def disclaimer():
    return render_template("disclaimer.html")

@app.route("/pdf-merge")
def pdf_merge_page():
    return render_template("pdf-merge.html")


@app.route("/merge-pdf", methods=["POST"])
def merge_pdf():

    files = request.files.getlist("pdfs")

    merger = PdfMerger()

    for file in files:

        filepath = os.path.join(
            UPLOAD_FOLDER,
            file.filename
        )

        file.save(filepath)

        merger.append(filepath)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "merged.pdf"
    )

    merger.write(output_path)
    merger.close()

    return send_file(
        output_path,
        as_attachment=True
    )



@app.route("/split-pdf", methods=["POST"])
def split_pdf():

    file = request.files["pdf"]
    page_number = int(request.form["page"])

    pdf_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(pdf_path)

    reader = PdfReader(pdf_path)

    writer = PdfWriter()

    writer.add_page(
        reader.pages[page_number - 1]
    )

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "split_page.pdf"
    )

    with open(output_path, "wb") as output_file:
        writer.write(output_file)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/pdf-compress")
def pdf_compress_page():
    return render_template("pdf-compress.html")
@app.route("/video-compress")
def video_compress_page():
    return render_template("video-compress.html")
@app.route("/video-converter")
def video_converter():
    return render_template("video-converter.html")
@app.route("/audio-converter")
def audio_converter():
    return render_template("audio-converter.html")


@app.route("/document-converter")
def document_converter():
    return render_template("document-converter.html")
@app.route("/convert-audio", methods=["POST"])
def convert_audio():
    file = request.files["audio"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    output_filename = f"{uuid.uuid4()}.{output_format}"

    output_path = os.path.join(
        CONVERTED_FOLDER,
        output_filename
    )

    subprocess.run([
        "ffmpeg",
        "-y",
        "-i",
        input_path,
        output_path
    ])

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/archive-converter")
def archive_converter():
    return render_template("archive-converter.html")

@app.route("/qr-generator")
def qr_generator():
    return render_template("qr-generator.html")


@app.route("/generate-qr", methods=["POST"])
def generate_qr():

    data = request.form["data"]

    qr = qrcode.make(data)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "qrcode.png"
    )

    qr.save(output_path)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/screenshot-tool")
def screenshot_tool():
    return render_template("screenshot-tool.html")


@app.route("/take-screenshot", methods=["POST"])
def take_screenshot():

    url = request.form["url"]

    options = webdriver.ChromeOptions()
    options.add_argument("--headless")

    driver = webdriver.Chrome(
        service=Service(
            ChromeDriverManager().install()
        ),
        options=options
    )

    driver.set_window_size(1920, 1080)

    driver.get(url)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "screenshot.png"
    )

    driver.save_screenshot(output_path)

    driver.quit()

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/url-converter")
def url_converter():
    return render_template("url-converter.html")


@app.route("/convert-url", methods=["POST"])
def convert_url():

    url = request.form["url"]

    response = requests.get(url)

    filename = url.split("/")[-1]

    if not filename:
        filename = "downloaded_file"

    output_path = os.path.join(
        CONVERTED_FOLDER,
        filename
    )

    with open(output_path, "wb") as f:
        f.write(response.content)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/gif-converter")
def gif_converter():
    return render_template("gif-converter.html")

@app.route("/convert-gif", methods=["POST"])
def convert_gif():

    file = request.files["gif"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        f"converted.{output_format}"
    )

    if output_format == "webp":

        image = Image.open(input_path)

        image.save(
            output_path,
            "WEBP"
        )

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/svg-converter")
def svg_converter():
    return render_template("svg-converter.html")

@app.route("/convert-svg", methods=["POST"])
def convert_svg():

    file = request.files["svg"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        "converted.png"
    )

    cairosvg.svg2png(
        url=input_path,
        write_to=output_path
    )

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/excel-converter")
def excel_converter():
    return render_template("excel-converter.html")

@app.route("/convert-excel", methods=["POST"])
def convert_excel():

    file = request.files["excel"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    df = pd.read_excel(input_path)

    if output_format == "csv":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.csv"
        )

        df.to_csv(output_path, index=False)

    elif output_format == "json":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.json"
        )

        df.to_json(
            output_path,
            orient="records"
        )

    elif output_format == "html":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.html"
        )

        df.to_html(
            output_path,
            index=False
        )

    elif output_format == "pdf":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.pdf"
        )

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=10)

        for row in df.values:
            pdf.cell(
                200,
                10,
                txt=" | ".join(map(str, row)),
                ln=True
            )

        pdf.output(output_path)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/powerpoint-converter")
def powerpoint_converter():
    return render_template("powerpoint-converter.html")

@app.route("/convert-powerpoint", methods=["POST"])
def convert_powerpoint():

    file = request.files["ppt"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    prs = Presentation(input_path)

    text_data = []

    for slide in prs.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text_data.append(shape.text)

    if output_format == "txt":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.txt"
        )

        with open(
            output_path,
            "w",
            encoding="utf-8"
        ) as f:

            f.write(
                "\n".join(text_data)
            )

    elif output_format == "json":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.json"
        )

        with open(
            output_path,
            "w",
            encoding="utf-8"
        ) as f:

            json.dump(
                text_data,
                f,
                indent=4
            )

    elif output_format == "pdf":

        subprocess.run([
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            "--headless",
            "--convert-to",
            "pdf",
            input_path,
            "--outdir",
            CONVERTED_FOLDER
        ])

        output_path = os.path.join(
            CONVERTED_FOLDER,
            os.path.splitext(
                os.path.basename(input_path)
            )[0] + ".pdf"
        )

    return send_file(
        output_path,
        as_attachment=True
    )

@app.route("/ebook-converter")
def ebook_converter():
    return render_template("ebook-converter.html")


@app.route("/convert-archive", methods=["POST"])
def convert_archive():

    file = request.files["archive"]
    action = request.form["action"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    # EXTRACT ZIP
    if action == "extract":

        extract_folder = os.path.join(
            CONVERTED_FOLDER,
            "extracted"
        )

        os.makedirs(
            extract_folder,
            exist_ok=True
        )

        import zipfile

        with zipfile.ZipFile(input_path, "r") as zip_ref:
            zip_ref.extractall(extract_folder)

        return "Archive Extracted Successfully"

    # ZIP → 7Z
    elif action == "zip_to_7z":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.7z"
        )

        import py7zr

        with py7zr.SevenZipFile(
            output_path,
            "w"
        ) as archive:

            archive.write(
                input_path,
                os.path.basename(input_path)
            )

        return send_file(
            output_path,
            as_attachment=True
        )

    # RAR → ZIP
    elif action == "rar_to_zip":

        import rarfile
        import zipfile

        temp_extract = os.path.join(
            CONVERTED_FOLDER,
            "rar_extract"
        )

        os.makedirs(
            temp_extract,
            exist_ok=True
        )

        with rarfile.RarFile(input_path) as rf:
            rf.extractall(temp_extract)

        output_path = os.path.join(
            CONVERTED_FOLDER,
            "converted.zip"
        )

        with zipfile.ZipFile(
            output_path,
            "w",
            zipfile.ZIP_DEFLATED
        ) as zipf:

            for root, dirs, files in os.walk(temp_extract):

                for f in files:

                    file_path = os.path.join(root, f)

                    zipf.write(
                        file_path,
                        os.path.relpath(
                            file_path,
                            temp_extract
                        )
                    )

        return send_file(
            output_path,
            as_attachment=True
        )

    return "Invalid Action"
@app.route("/convert-ebook", methods=["POST"])
def convert_ebook():

    file = request.files["ebook"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        secure_filename(file.filename)
    )

    file.save(input_path)

    output_path = os.path.join(
        CONVERTED_FOLDER,
        f"converted.{output_format}"
    )

    subprocess.run([
        r"C:\Program Files\Calibre2\ebook-convert.exe",
        input_path,
        output_path
    ])

    return send_file(
        output_path,
        as_attachment=True
    )


 
@app.route("/convert-document", methods=["POST"])
def convert_document():

    file = request.files["document"]
    output_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    if output_format == "pdf":

        output_path = os.path.join(
            CONVERTED_FOLDER,
            f"{uuid.uuid4()}.pdf"
        )

        pdf = canvas.Canvas(output_path)

        with open(input_path, "r", encoding="utf-8") as f:
            text = f.read()

        pdf.drawString(50, 800, text[:1000])
        pdf.save()

        return send_file(
            output_path,
            as_attachment=True
        )

    return "Conversion Coming Soon"
@app.route("/compress-video", methods=["POST"])
def compress_video():

    file = request.files["video"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    output_path = os.path.join(
    CONVERTED_FOLDER,
    f"{uuid.uuid4()}.mp4"
)

    subprocess.run([
    "ffmpeg",
    "-y",
    "-i", input_path,
    "-vcodec", "libx264",
    "-crf", "30",
    output_path
], check=True)
    print("FFMPEG DONE")


    before_size = os.path.getsize(input_path) / (1024 * 1024)
    after_size = os.path.getsize(output_path) / (1024 * 1024)

    print("================================")
    print("Original Size :", round(before_size, 2), "MB")
    print("Compressed Size :", round(after_size, 2), "MB")
    print("Saved :", round(before_size - after_size, 2), "MB")
    print("================================")

    return send_file(
        output_path,
        as_attachment=True
    )
    
@app.route("/convert-video", methods=["POST"])
def convert_video():

    file = request.files["video"]
    target_format = request.form["format"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    output_filename = f"{uuid.uuid4()}.{target_format}"
    output_path = os.path.join(
        CONVERTED_FOLDER,
        output_filename
    )

    subprocess.run([
        "ffmpeg",
        "-y",
        "-i",
        input_path,
        output_path
    ])

    return send_file(
        output_path,
        as_attachment=True
    )

@app.route("/compress-pdf", methods=["POST"])
def compress_pdf():




    file = request.files["pdf"]

    input_path = os.path.join(
        UPLOAD_FOLDER,
        file.filename
    )

    file.save(input_path)

    output_path = os.path.join(
    CONVERTED_FOLDER,
    f"{uuid.uuid4()}.mp4"
)
    pdf = pikepdf.open(input_path)
    pdf.save(output_path)

    return send_file(
        output_path,
        as_attachment=True
    )
@app.route("/video")
def video():
    return "VIDEO OK"
print("APP LOADED")
print(app.url_map)
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)